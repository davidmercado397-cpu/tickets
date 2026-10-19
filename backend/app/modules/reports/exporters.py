"""Exportación de reportes a Excel (openpyxl) y PDF (fpdf2)."""

from __future__ import annotations

import io
from datetime import datetime

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.documents.public import latin1


def to_xlsx(title: str, header: list[str], rows: list[list], color: str = "#0B3D91") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    fill = PatternFill("solid", fgColor=color.lstrip("#"))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, _ in enumerate(header, start=1):
        width = max(len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, min(ws.max_row, 200) + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary_pdf(tenant: dict, summary: dict, period: str) -> bytes:
    pdf = FPDF(orientation="P", format="Letter")
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, latin1(f"Informe de PQRS - {tenant['name']}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 5, latin1(f"Periodo: {period}  ·  Generado: {datetime.now():%Y-%m-%d %H:%M}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    kpis = [
        ("Total", summary["total"]), ("Abiertas", summary["open"]), ("Respondidas", summary["responded"]),
        ("Cumplimiento", f"{summary['compliance']}%" if summary["compliance"] is not None else "—"),
        ("Vencidas", summary["sla"]["overdue"]),
    ]
    pdf.set_font("Helvetica", "B", 10)
    w = (pdf.w - 30) / len(kpis)
    for label, _ in kpis:
        pdf.cell(w, 6, latin1(label), border=1, align="C")
    pdf.ln()
    pdf.set_font("Helvetica", "", 11)
    for _, value in kpis:
        pdf.cell(w, 8, latin1(str(value)), border=1, align="C")
    pdf.ln(12)
    for key, title in (("by_type", "Por tipo"), ("by_area", "Por área"), ("by_category", "Por categoría"),
                       ("by_channel", "Por canal"), ("by_client", "Por cliente"), ("by_assignee", "Por responsable")):
        groups = summary.get(key) or []
        if not groups:
            continue
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 7, latin1(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "B", 8.5)
        cols = [("Nombre", 80), ("Casos", 20), ("Abiertas", 20), ("Vencidas", 20), ("Cumplidas", 20), ("% cumpl.", 26)]
        for label, cw in cols:
            pdf.cell(cw, 6, latin1(label), border=1, fill=False)
        pdf.ln()
        pdf.set_font("Helvetica", "", 8.5)
        for g in groups[:20]:
            values = [g["name"], g["count"], g["open"], g["overdue"], g["met"],
                      f"{g['compliance']}%" if g["compliance"] is not None else "—"]
            for (_, cw), v in zip(cols, values):
                pdf.cell(cw, 5.5, latin1(str(v))[:48], border=1)
            pdf.ln()
        pdf.ln(4)
    return bytes(pdf.output())
