"""Importación masiva de clientes, puestos y metas de visitas desde Excel (.xlsx).

Hojas: "Clientes" y "Puestos" (también se acepta "Sedes"). Las columnas se reconocen por su título, sin
importar mayúsculas ni tildes. Todo o nada: si hay cualquier error no se guarda ningún cambio, y con
`dry_run` solo se valida y se informa qué pasaría."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.context import Ctx
from app.core.errors import ValidationError
from app.core.text import slugify
from app.modules.clients.models import Client, Site
from app.modules.clients.service import normalize_nit

MAX_ROWS = 5000
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# (clave, título en la plantilla, obligatorio, alias aceptados)
CLIENT_COLUMNS = [
    ("nit", "NIT", True, ["nit", "nit cliente"]),
    ("name", "Razón social", True, ["razon social", "nombre", "cliente", "nombre cliente"]),
    ("contact_name", "Contacto", False, ["contacto", "nombre contacto"]),
    ("contact_email", "Correo contacto", False, ["correo contacto", "correo", "email"]),
    ("contact_phone", "Teléfono", False, ["telefono", "celular", "telefono contacto"]),
    ("address", "Dirección", False, ["direccion"]),
    ("city", "Ciudad", False, ["ciudad"]),
    ("contract_start", "Inicio contrato", False, ["inicio contrato", "fecha inicio"]),
    ("contract_end", "Fin contrato", False, ["fin contrato", "fecha fin"]),
    ("quota", "Meta visitas/mes (cliente)", False, ["meta visitas/mes (cliente)", "meta visitas", "meta visitas mes", "visitas mes"]),
    ("notes", "Notas", False, ["notas", "observaciones"]),
]
SITE_COLUMNS = [
    ("nit", "NIT cliente", True, ["nit cliente", "nit"]),
    ("name", "Nombre puesto", True, ["nombre puesto", "puesto", "sede", "nombre sede", "nombre"]),
    ("code", "Código", False, ["codigo", "codigo puesto"]),
    ("address", "Dirección", False, ["direccion"]),
    ("city", "Ciudad", False, ["ciudad"]),
    ("contact_name", "Contacto", False, ["contacto"]),
    ("contact_phone", "Teléfono", False, ["telefono", "celular"]),
    ("latitude", "Latitud", False, ["latitud", "lat"]),
    ("longitude", "Longitud", False, ["longitud", "lng", "lon"]),
    ("quota", "Meta visitas/mes (puesto)", False, ["meta visitas/mes (puesto)", "meta visitas", "meta visitas mes", "visitas mes"]),
]


def _norm(text: Any) -> str:
    s = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


# --- Plantilla ---

def build_template() -> bytes:
    wb = Workbook()
    head = PatternFill("solid", fgColor="0B3D91")
    examples = {
        "Clientes": (CLIENT_COLUMNS, [
            ["900123456-7", "Conjunto Residencial Los Robles", "María Gómez", "admin@losrobles.co", "3001234567",
             "Cra 10 # 20-30", "Santiago de Cali", "2026-01-01", "2026-12-31", 2, "Cliente desde 2020"],
            ["800987654-1", "Centro Comercial Plaza Norte", "Jorge Ruiz", "seguridad@plazanorte.co", "3109876543",
             "Av 6N # 40-10", "Santiago de Cali", "2025-06-01", "", "", ""],
        ]),
        "Puestos": (SITE_COLUMNS, [
            ["900123456-7", "Portería principal", "ROBLES-01", "Cra 10 # 20-30", "Santiago de Cali", "Carlos Pérez",
             "3001112233", 3.4516, -76.532, 4],
            ["900123456-7", "Portería vehicular", "ROBLES-02", "Calle 21 # 10-05", "Santiago de Cali", "", "", "", "", 2],
            ["800987654-1", "Parqueadero sótano", "", "Av 6N # 40-10", "Santiago de Cali", "", "", "", "", 3],
        ]),
    }
    first = True
    for title, (cols, rows) in examples.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append([f"{c[1]}{' *' if c[2] else ''}" for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(c[1]) + 6)
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("Instrucciones")
    lines = [
        "Importación de clientes, puestos y metas de visitas",
        "",
        "1. Hoja Clientes: una fila por cliente. Se identifica por el NIT: si ya existe se actualiza, si no se crea.",
        "2. Hoja Puestos: una fila por puesto o sede. 'NIT cliente' debe existir en la hoja Clientes o en la plataforma.",
        "   El puesto se identifica por su Código; si no tiene código, por cliente + nombre del puesto.",
        "3. Meta visitas/mes: cuántas visitas de supervisión se exigen al mes.",
        "   - En Clientes: cuenta las visitas a cualquier puesto del cliente.",
        "   - En Puestos: cuenta solo las visitas a ese puesto.",
        "   - Vacío = no cambia la meta actual · 0 = sin meta.",
        "4. Fechas en formato AAAA-MM-DD o DD/MM/AAAA. Las columnas con * son obligatorias.",
        "5. Antes de guardar, la plataforma valida todo el archivo: si hay errores no se guarda nada.",
        "6. Las filas de ejemplo se pueden borrar o reemplazar.",
    ]
    for line in lines:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Lectura ---

@dataclass
class Report:
    clients_created: int = 0
    clients_updated: int = 0
    sites_created: int = 0
    sites_updated: int = 0
    quotas_set: int = 0
    errors: list[dict] = field(default_factory=list)

    def error(self, sheet: str, row: int, message: str) -> None:
        if len(self.errors) < 200:
            self.errors.append({"sheet": sheet, "row": row, "message": message})


def _sheet(wb, *names: str):
    for ws in wb.worksheets:
        if _norm(ws.title) in names:
            return ws
    return None


def _rows(ws, columns, sheet: str, report: Report) -> list[tuple[int, dict]]:
    header = [_norm(c.value).rstrip(" *") for c in ws[1]]
    index: dict[str, int] = {}
    for key, title, _required, aliases in columns:
        names = {_norm(title), *aliases}
        for i, h in enumerate(header):
            if h in names and key not in index:
                index[key] = i
    missing = [title for key, title, required, _ in columns if required and key not in index]
    if missing:
        report.error(sheet, 1, "Faltan columnas obligatorias: " + ", ".join(missing))
        return []
    out = []
    for n, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if n - 1 > MAX_ROWS:
            report.error(sheet, n, f"Máximo {MAX_ROWS} filas por hoja.")
            break
        values = {k: (raw[i] if i < len(raw) else None) for k, i in index.items()}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        out.append((n, values))
    return out


def _text(v: Any, limit: int) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s[:limit] if s else None


def _date(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"fecha inválida '{s}' (usa AAAA-MM-DD o DD/MM/AAAA)")


def _int(v: Any) -> int | None:
    if v in (None, ""):
        return None
    try:
        n = float(str(v).replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"'{v}' no es un número") from exc
    if not n.is_integer() or n < 0 or n > 500:
        raise ValueError("la meta debe ser un entero entre 0 y 500")
    return int(n)


def _float(v: Any, lo: float, hi: float, name: str) -> float | None:
    if v in (None, ""):
        return None
    try:
        n = float(str(v).replace(",", "."))
    except ValueError as exc:
        raise ValueError(f"{name} inválida") from exc
    if not lo <= n <= hi:
        raise ValueError(f"{name} fuera de rango")
    return n


# --- Importación ---

async def import_workbook(session: AsyncSession, ctx: Ctx, content: bytes, *, dry_run: bool) -> dict:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - archivo corrupto o no es Excel
        raise ValidationError("No se pudo leer el archivo. Usa la plantilla en formato .xlsx.") from exc
    ws_clients = _sheet(wb, "clientes")
    ws_sites = _sheet(wb, "puestos", "sedes")
    if ws_clients is None and ws_sites is None:
        raise ValidationError("El archivo debe tener una hoja 'Clientes' y/o una hoja 'Puestos'.")

    report = Report()
    can_quota = ctx.can("visits.manage")
    tid = ctx.tid
    existing = {c.nit: c for c in (await session.execute(select(Client).where(Client.tenant_id == tid))).scalars()}
    quotas: list[tuple[str, int, int, Client | Site]] = []  # (hoja, fila, meta, cliente o puesto)

    # Clientes
    seen_nits: dict[str, int] = {}
    for n, row in _rows(ws_clients, CLIENT_COLUMNS, "Clientes", report) if ws_clients else []:
        nit = normalize_nit(_text(row.get("nit"), 30) or "")
        name = _text(row.get("name"), 200)
        if len(nit) < 3:
            report.error("Clientes", n, "NIT obligatorio.")
            continue
        if not name or len(name) < 2:
            report.error("Clientes", n, "Razón social obligatoria.")
            continue
        if nit in seen_nits:
            report.error("Clientes", n, f"NIT repetido (ya está en la fila {seen_nits[nit]}).")
            continue
        seen_nits[nit] = n
        try:
            email = _text(row.get("contact_email"), 200)
            if email and not EMAIL.match(email):
                raise ValueError(f"correo inválido '{email}'")
            values = {
                "name": name,
                "contact_name": _text(row.get("contact_name"), 200),
                "contact_email": email,
                "contact_phone": _text(row.get("contact_phone"), 50),
                "address": _text(row.get("address"), 250),
                "city": _text(row.get("city"), 100),
                "contract_start": _date(row.get("contract_start")),
                "contract_end": _date(row.get("contract_end")),
                "notes": _text(row.get("notes"), 5000),
            }
            quota = _int(row.get("quota"))
        except ValueError as exc:
            report.error("Clientes", n, str(exc).capitalize() + ".")
            continue
        client = existing.get(nit)
        if client is None:
            client = Client(tenant_id=tid, nit=nit, active=True, **values)
            session.add(client)
            existing[nit] = client
            report.clients_created += 1
        else:
            # Las celdas vacías no borran datos existentes.
            for k, v in values.items():
                if v is not None:
                    setattr(client, k, v)
            report.clients_updated += 1
        if quota is not None:
            quotas.append(("Clientes", n, quota, client))
    await session.flush()

    # Puestos
    sites = list((await session.execute(select(Site).where(Site.tenant_id == tid))).scalars())
    by_code = {s.code.lower(): s for s in sites}
    by_name = {(s.client_id, s.name.strip().lower()): s for s in sites}
    seen_sites: dict[Any, int] = {}
    for n, row in _rows(ws_sites, SITE_COLUMNS, "Puestos", report) if ws_sites else []:
        nit = normalize_nit(_text(row.get("nit"), 30) or "")
        name = _text(row.get("name"), 200)
        client = existing.get(nit)
        if client is None:
            report.error("Puestos", n, f"No existe un cliente con NIT '{nit or '(vacío)'}' en la hoja Clientes ni en la plataforma.")
            continue
        if not name or len(name) < 2:
            report.error("Puestos", n, "Nombre del puesto obligatorio.")
            continue
        code = _text(row.get("code"), 40)
        if code and not re.fullmatch(r"[A-Za-z0-9_-]{2,40}", code):
            report.error("Puestos", n, "El código solo admite letras, números, guion y guion bajo (2 a 40).")
            continue
        key = ("code", code.lower()) if code else ("name", client.id, name.strip().lower())
        if key in seen_sites:
            report.error("Puestos", n, f"Puesto repetido (ya está en la fila {seen_sites[key]}).")
            continue
        seen_sites[key] = n
        try:
            values = {
                "name": name,
                "address": _text(row.get("address"), 250),
                "city": _text(row.get("city"), 100),
                "contact_name": _text(row.get("contact_name"), 200),
                "contact_phone": _text(row.get("contact_phone"), 50),
                "latitude": _float(row.get("latitude"), -90, 90, "Latitud"),
                "longitude": _float(row.get("longitude"), -180, 180, "Longitud"),
            }
            quota = _int(row.get("quota"))
        except ValueError as exc:
            report.error("Puestos", n, str(exc).capitalize() + ".")
            continue
        site = by_code.get(code.lower()) if code else by_name.get((client.id, name.strip().lower()))
        if site is not None and site.client_id != client.id:
            report.error("Puestos", n, f"El código '{code}' ya pertenece a un puesto de otro cliente.")
            continue
        if site is None:
            if not code:
                base, i = slugify(name, 36, "-"), 1
                code = base
                while code.lower() in by_code:
                    i += 1
                    code = f"{base[:36]}-{i}"
            site = Site(tenant_id=tid, client_id=client.id, code=code, active=True, **values)
            session.add(site)
            by_code[code.lower()] = site
            by_name[(client.id, name.strip().lower())] = site
            report.sites_created += 1
        else:
            for k, v in values.items():
                if v is not None:
                    setattr(site, k, v)
            site.active = True
            report.sites_updated += 1
        if quota is not None:
            quotas.append(("Puestos", n, quota, site))
    await session.flush()

    # Metas de visitas (las gestiona el módulo de visitas).
    if quotas and not can_quota:
        report.error(quotas[0][0], quotas[0][1],
                     "Tu rol no puede configurar metas de visitas (permiso de gestionar visitas). Deja la columna vacía.")
    elif quotas:
        from app.modules.visits.public import set_quota_count  # import diferido: visitas depende de clientes

        for _sheet_name, _row, n_visits, target in quotas:
            if isinstance(target, Site):
                changed = await set_quota_count(session, tid, site_id=target.id, visits_per_month=n_visits)
            else:
                changed = await set_quota_count(session, tid, client_id=target.id, visits_per_month=n_visits)
            report.quotas_set += int(changed)

    ok = not report.errors
    if ok and not dry_run:
        await session.commit()
    else:
        await session.rollback()
    return {
        "dry_run": dry_run,
        "saved": ok and not dry_run,
        "clients_created": report.clients_created,
        "clients_updated": report.clients_updated,
        "sites_created": report.sites_created,
        "sites_updated": report.sites_updated,
        "quotas_set": report.quotas_set,
        "errors": report.errors,
    }

