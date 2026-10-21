"""Importación de clientes, puestos y metas de visitas desde Excel."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from tests.conftest import login

CLIENT_HEAD = ["NIT *", "Razón social *", "Contacto", "Correo contacto", "Ciudad", "Inicio contrato",
               "Meta visitas/mes (cliente)"]
SITE_HEAD = ["NIT cliente *", "Nombre puesto *", "Código", "Dirección", "Latitud", "Longitud",
             "Meta visitas/mes (puesto)"]


def _xlsx(clients: list[list], sites: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(CLIENT_HEAD)
    for r in clients:
        ws.append(r)
    ws2 = wb.create_sheet("Puestos")
    ws2.append(SITE_HEAD)
    for r in sites:
        ws2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, h, content: bytes, dry_run: bool):
    return client.post("/api/clients/import", headers=h, data={"dry_run": str(dry_run).lower()},
                       files={"file": ("clientes.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


GOOD = _xlsx(
    [["900.123.456-7", "Conjunto Los Robles", "María", "maria@robles.co", "Cali", "2026-01-01", 2],
     ["800987654-1", "CC Plaza Norte", None, None, "Cali", "15/06/2025", None]],
    [["900123456-7", "Portería principal", "ROBLES-01", "Cra 10", 3.45, -76.53, 4],
     ["900123456-7", "Portería vehicular", None, "Calle 21", None, None, 2],
     ["800987654-1", "Parqueadero", None, None, None, None, None]],
)


def test_template_has_expected_sheets(client, admin):
    r = client.get("/api/clients/import-template", headers=admin)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Clientes", "Puestos", "Instrucciones"]
    assert wb["Puestos"]["J1"].value == "Meta visitas/mes (puesto)"


def test_dry_run_validates_without_saving_then_imports(client, admin):
    r = _upload(client, admin, GOOD, dry_run=True)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["errors"] == [] and body["saved"] is False
    assert (body["clients_created"], body["sites_created"], body["quotas_set"]) == (2, 3, 3)
    assert client.get("/api/clients", headers=admin).json() == []  # nada guardado

    body = _upload(client, admin, GOOD, dry_run=False).json()
    assert body["saved"] is True
    clients = {c["nit"]: c for c in client.get("/api/clients", headers=admin).json()}
    robles = clients["900123456-7"]
    assert robles["contract_start"] == "2026-01-01" and len(robles["sites"]) == 2
    assert clients["800987654-1"]["contract_start"] == "2025-06-15"
    quotas = client.get("/api/visits/quotas", headers=admin).json()
    by_target = {(q["client"] or q["site"])["name"]: q["visits_per_month"] for q in quotas}
    assert by_target == {"Conjunto Los Robles": 2, "Portería principal": 4, "Portería vehicular": 2}

    # Reimportar actualiza (no duplica) y las celdas vacías no borran datos.
    again = _xlsx([["900123456-7", "Conjunto Los Robles II", None, None, None, None, 3]],
                  [["900123456-7", "Portería principal", "ROBLES-01", None, None, None, 0]])
    body = _upload(client, admin, again, dry_run=False).json()
    assert (body["clients_created"], body["clients_updated"], body["sites_updated"]) == (0, 1, 1)
    robles = next(c for c in client.get("/api/clients", headers=admin).json() if c["nit"] == "900123456-7")
    assert robles["name"] == "Conjunto Los Robles II" and robles["contact_email"] == "maria@robles.co"
    assert len(robles["sites"]) == 2
    quotas = {(q["client"] or q["site"])["name"]: q["visits_per_month"]
              for q in client.get("/api/visits/quotas", headers=admin).json()}
    assert quotas["Conjunto Los Robles II"] == 3 and quotas["Portería principal"] == 0


def test_any_error_saves_nothing_and_reports_rows(client, admin):
    bad = _xlsx(
        [["900123456-7", "Cliente bueno", None, "correo-malo", None, None, None],
         ["900123456-7", "Repetido", None, None, None, None, None],
         ["", "Sin NIT", None, None, None, None, None]],
        [["111", "Puesto huérfano", None, None, None, None, None],
         ["900123456-7", "Puesto", None, None, 120, None, "dos"]],
    )
    body = _upload(client, admin, bad, dry_run=False).json()
    assert body["saved"] is False
    errors = {(e["sheet"], e["row"]) for e in body["errors"]}
    assert {("Clientes", 2), ("Clientes", 3), ("Clientes", 4), ("Puestos", 2), ("Puestos", 3)} <= errors
    assert client.get("/api/clients", headers=admin).json() == []


def test_missing_columns_and_non_excel(client, admin):
    wb = Workbook()
    wb.active.title = "Clientes"
    wb.active.append(["Nombre"])
    buf = io.BytesIO()
    wb.save(buf)
    body = _upload(client, admin, buf.getvalue(), dry_run=True).json()
    assert "Faltan columnas obligatorias" in body["errors"][0]["message"]
    r = _upload(client, admin, b"no soy un excel", dry_run=True)
    assert r.status_code == 422


def test_quotas_require_visits_permission(client, admin):
    roles = {r["code"]: r["id"] for r in client.get("/api/roles", headers=admin).json()}
    r = client.post("/api/roles", headers=admin, json={"name": "Comercial", "permissions": ["clients.manage"]})
    client.post("/api/users", headers=admin, json={"email": "com@demo.nextbyte.com.co", "full_name": "Comercial",
                                                   "password": "comercial-123", "role_ids": [r.json()["id"]]})
    assert "admin" in roles
    h = login(client, "com@demo.nextbyte.com.co", "comercial-123")
    body = _upload(client, h, GOOD, dry_run=True).json()
    assert body["errors"] and "metas de visitas" in body["errors"][0]["message"]
    no_quota = _xlsx([["900123456-7", "Sin meta", None, None, None, None, None]], [])
    assert _upload(client, h, no_quota, dry_run=False).json()["saved"] is True


def test_client_nit_can_be_corrected(client, admin):
    """Clientes cargados con NIT provisional: el NIT se corrige después, sin duplicar otro."""
    a = client.post("/api/clients", headers=admin, json={"nit": "0000000001", "name": "Provisional A"}).json()
    client.post("/api/clients", headers=admin, json={"nit": "0000000002", "name": "Provisional B"})
    r = client.patch(f"/api/clients/{a['id']}", headers=admin, json={"nit": "890.300.123-4", "name": "Provisional A"})
    assert r.status_code == 200, r.text
    assert r.json()["nit"] == "890300123-4"
    r = client.patch(f"/api/clients/{a['id']}", headers=admin, json={"nit": "0000000002"})
    assert r.status_code == 409
